"""Builds the three spreadsheet templates in the kit."""

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

BRAND = "4F46E5"
BRAND_LIGHT = "EEF2FF"
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor=BRAND)
INPUT_FILL = PatternFill("solid", fgColor="FFFBEB")
TOTAL_FILL = PatternFill("solid", fgColor=BRAND_LIGHT)
RED_FILL = PatternFill("solid", fgColor="FEE2E2")
GREEN_FILL = PatternFill("solid", fgColor="DCFCE7")
THIN = Side(style="thin", color="E5E7EB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TITLE_FONT = Font(bold=True, size=16, color=BRAND)
MONEY = '#,##0.00;[Red]-#,##0.00'


def _title(ws, text, sub):
    ws["A1"] = text
    ws["A1"].font = TITLE_FONT
    ws["A2"] = sub
    ws["A2"].font = Font(italic=True, color="6B7280")


def _header(ws, row, headers, widths):
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 24
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def _list_validation(ws, formula, rng):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(rng)


# --------------------------------------------------------------------------
# 1. Profit & Cash Flow Tracker
# --------------------------------------------------------------------------
INCOME_CATS = ["Product sales", "Services", "Retainers", "Affiliate", "Other income"]
EXPENSE_CATS = ["Software & tools", "Advertising", "Contractors", "Equipment", "Education",
                "Fees & banking", "Travel", "Office", "Other expense"]
TX_ROWS = 1000
CAT_LAST = 30


def build_profit_tracker(path):
    wb = Workbook()

    # Settings
    st = wb.active
    st.title = "Settings"
    _title(st, "Settings", "Yellow cells are yours to edit.")
    rows = [("Business name", "My Business"), ("Year", 2026), ("Currency symbol", "$"),
            ("Tax set-aside %", 0.25), ("Monthly revenue goal", 5000)]
    for i, (k, v) in enumerate(rows, start=4):
        st.cell(row=i, column=1, value=k).font = Font(bold=True)
        c = st.cell(row=i, column=2, value=v)
        c.fill = INPUT_FILL
        c.border = BORDER
    st["B7"].number_format = "0%"
    st["B8"].number_format = MONEY
    st.column_dimensions["A"].width = 24
    st.column_dimensions["B"].width = 20

    st["D3"] = "Category"
    st["E3"] = "Type"
    st["D3"].font = st["E3"].font = Font(bold=True)
    cats = [(c, "Income") for c in INCOME_CATS] + [(c, "Expense") for c in EXPENSE_CATS]
    for i, (cat, kind) in enumerate(cats, start=4):
        st.cell(row=i, column=4, value=cat).fill = INPUT_FILL
        st.cell(row=i, column=5, value=kind)
    for r in range(4 + len(cats), CAT_LAST + 1):
        st.cell(row=r, column=4).fill = INPUT_FILL
    st.column_dimensions["D"].width = 22
    st.column_dimensions["E"].width = 12
    st["G4"] = f"Tip: rename categories or add new ones (down to row {CAT_LAST}) and the dropdowns update."
    st["G4"].font = Font(italic=True, color="6B7280")

    # Transactions
    tx = wb.create_sheet("Transactions")
    _title(tx, "Transactions", "One row per sale or expense. Enter amounts as positive numbers.")
    _header(tx, 4, ["Date", "Type", "Category", "Description", "Client / Vendor", "Amount"],
            [13, 11, 20, 36, 22, 14])
    last = 4 + TX_ROWS
    _list_validation(tx, '"Income,Expense"', f"B5:B{last}")
    _list_validation(tx, f"=Settings!$D$4:$D${CAT_LAST}", f"C5:C{last}")
    for r in range(5, last + 1):
        tx.cell(row=r, column=1).number_format = "yyyy-mm-dd"
        tx.cell(row=r, column=6).number_format = MONEY
    tx.conditional_formatting.add(f"B5:B{last}", CellIsRule(operator="equal", formula=['"Income"'], fill=GREEN_FILL))
    tx.conditional_formatting.add(f"B5:B{last}", CellIsRule(operator="equal", formula=['"Expense"'], fill=RED_FILL))
    samples = [
        ("2026-01-05", "Income", "Services", "Website audit", "Acme Co", 750),
        ("2026-01-09", "Expense", "Software & tools", "Email platform", "EmailCo", 29),
        ("2026-01-15", "Income", "Product sales", "Template sales", "Online store", 420),
        ("2026-01-20", "Expense", "Advertising", "Social ads test", "Ads", 120),
    ]
    from datetime import date
    for i, (d, t, c, desc, who, amt) in enumerate(samples, start=5):
        tx.cell(row=i, column=1, value=date.fromisoformat(d))
        tx.cell(row=i, column=2, value=t)
        tx.cell(row=i, column=3, value=c)
        tx.cell(row=i, column=4, value=desc)
        tx.cell(row=i, column=5, value=who)
        tx.cell(row=i, column=6, value=amt)
    tx["H4"] = "Sample rows: delete them when you start."
    tx["H4"].font = Font(italic=True, color="6B7280")
    tx.auto_filter.ref = f"A4:F{last}"

    # Dashboard
    db = wb.create_sheet("Dashboard", 0)
    _title(db, "Profit & Cash Flow Dashboard", "Fills in automatically from Transactions. Nothing to edit here.")
    _header(db, 4, ["Month", "Income", "Expenses", "Profit", "Margin", "Tax set-aside",
                    "Take-home", "vs Goal", "Running cash"], [12, 14, 14, 14, 10, 14, 14, 12, 15])
    db.freeze_panes = None
    months = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
    rng = f"Transactions!$F$5:$F${last}"
    typ = f"Transactions!$B$5:$B${last}"
    yr_check = f"Transactions!$A$5:$A${last}"
    for i, m in enumerate(months, start=1):
        r = 4 + i
        db.cell(row=r, column=1, value=m).font = Font(bold=True)
        # Filter by year using a date window so the tracker can be reused across years.
        start = f"DATE(Settings!$B$5,{i},1)"
        end = f"DATE(Settings!$B$5,{i}+1,1)"
        db.cell(row=r, column=2, value=f'=SUMIFS({rng},{typ},"Income",{yr_check},">="&{start},{yr_check},"<"&{end})')
        db.cell(row=r, column=3, value=f'=SUMIFS({rng},{typ},"Expense",{yr_check},">="&{start},{yr_check},"<"&{end})')
        db.cell(row=r, column=4, value=f"=B{r}-C{r}")
        db.cell(row=r, column=5, value=f'=IF(B{r}=0,"",D{r}/B{r})')
        db.cell(row=r, column=6, value=f"=MAX(0,D{r})*Settings!$B$7")
        db.cell(row=r, column=7, value=f"=D{r}-F{r}")
        db.cell(row=r, column=8, value=f'=IF(Settings!$B$8=0,"",B{r}/Settings!$B$8)')
        db.cell(row=r, column=9, value=f"=G{r}" if i == 1 else f"=I{r - 1}+G{r}")
        for col in (2, 3, 4, 6, 7, 9):
            db.cell(row=r, column=col).number_format = MONEY
        db.cell(row=r, column=5).number_format = "0%"
        db.cell(row=r, column=8).number_format = "0%"
        for col in range(1, 10):
            db.cell(row=r, column=col).border = BORDER
    tr = 17
    db.cell(row=tr, column=1, value="Year total").font = Font(bold=True)
    for col in (2, 3, 4, 6, 7):
        L = get_column_letter(col)
        c = db.cell(row=tr, column=col, value=f"=SUM({L}5:{L}16)")
        c.number_format = MONEY
    db.cell(row=tr, column=5, value="=IF(B17=0,\"\",D17/B17)").number_format = "0%"
    for col in range(1, 10):
        db.cell(row=tr, column=col).fill = TOTAL_FILL
        db.cell(row=tr, column=col).font = Font(bold=True)
        db.cell(row=tr, column=col).border = BORDER
    db.conditional_formatting.add("D5:D17", CellIsRule(operator="lessThan", formula=["0"], fill=RED_FILL))
    db.conditional_formatting.add("H5:H16", CellIsRule(operator="greaterThanOrEqual", formula=["1"], fill=GREEN_FILL))

    db["A20"] = "Top expense categories (year)"
    db["A20"].font = Font(bold=True, color=BRAND)
    for i in range(len(EXPENSE_CATS)):
        r = 21 + i
        db.cell(row=r, column=1, value=f"=Settings!D{4 + len(INCOME_CATS) + i}")
        db.cell(row=r, column=2, value=(
            f'=SUMIFS({rng},{typ},"Expense",Transactions!$C$5:$C${last},A{r},'
            f'{yr_check},">="&DATE(Settings!$B$5,1,1),{yr_check},"<"&DATE(Settings!$B$5+1,1,1))'
        )).number_format = MONEY
    db.column_dimensions["A"].width = 22

    wb.active = 0
    wb.save(path)


# --------------------------------------------------------------------------
# 2. Content Calendar
# --------------------------------------------------------------------------
PLATFORMS = ["Instagram", "LinkedIn", "TikTok", "YouTube", "X / Twitter", "Facebook", "Newsletter", "Blog", "Pinterest"]
PILLARS = ["Educational", "Story", "Proof / results", "Offer", "Behind the scenes"]
FORMATS = ["Carousel", "Short video", "Long video", "Text post", "Thread", "Image", "Article", "Email", "Live"]
STATUSES = ["Idea", "Drafting", "Ready", "Scheduled", "Published"]
CAL_ROWS = 365


def build_content_calendar(path):
    from datetime import date, timedelta
    wb = Workbook()
    ws = wb.active
    ws.title = "Calendar"
    _title(ws, "Content Calendar", "Paste your AI-generated 30-day plan here (prompt 2.1). Status turns green when published.")
    _header(ws, 4, ["Date", "Day", "Platform", "Pillar", "Hook / Title", "Format", "CTA", "Status", "Link", "Notes"],
            [12, 6, 14, 17, 46, 13, 24, 12, 26, 30])
    last = 4 + CAL_ROWS

    lists = wb.create_sheet("Lists")
    for col, (name, items) in enumerate([("Platforms", PLATFORMS), ("Pillars", PILLARS),
                                         ("Formats", FORMATS), ("Statuses", STATUSES)], start=1):
        lists.cell(row=1, column=col, value=name).font = Font(bold=True)
        for i, it in enumerate(items, start=2):
            lists.cell(row=i, column=col, value=it)
        lists.column_dimensions[get_column_letter(col)].width = 18
    _list_validation(ws, "=Lists!$A$2:$A$20", f"C5:C{last}")
    _list_validation(ws, "=Lists!$B$2:$B$20", f"D5:D{last}")
    _list_validation(ws, "=Lists!$C$2:$C$20", f"F5:F{last}")
    _list_validation(ws, "=Lists!$D$2:$D$20", f"H5:H{last}")

    start = date(2026, 1, 1)
    for i in range(CAL_ROWS):
        r = 5 + i
        ws.cell(row=r, column=1, value=start + timedelta(days=i)).number_format = "yyyy-mm-dd"
        ws.cell(row=r, column=2, value=f'=IF(A{r}="","",TEXT(A{r},"ddd"))')
        ws.cell(row=r, column=5).alignment = Alignment(wrap_text=True)
    ws["L4"] = "Tip: change A5 to your start date and drag down to reset the dates."
    ws["L4"].font = Font(italic=True, color="6B7280")
    ws.conditional_formatting.add(f"A5:J{last}", FormulaRule(formula=['$H5="Published"'], fill=GREEN_FILL))
    ws.conditional_formatting.add(f"A5:J{last}", FormulaRule(formula=['AND($A5=TODAY(),$H5<>"Published")'], fill=INPUT_FILL))
    ws.auto_filter.ref = f"A4:J{last}"

    sm = wb.create_sheet("Summary", 0)
    _title(sm, "Content Pipeline Summary", "Counts update automatically.")
    _header(sm, 4, ["Status", "Posts", "", "Platform", "Planned", "Published"], [22, 10, 8, 16, 10, 11])
    sm.freeze_panes = None
    sm["C4"].fill = PatternFill()
    for i, s in enumerate(STATUSES, start=5):
        sm.cell(row=i, column=1, value=s)
        sm.cell(row=i, column=2, value=f'=COUNTIF(Calendar!$H$5:$H${last},A{i})')
    for i, p in enumerate(PLATFORMS, start=5):
        sm.cell(row=i, column=4, value=p)
        sm.cell(row=i, column=5, value=f'=COUNTIF(Calendar!$C$5:$C${last},D{i})')
        sm.cell(row=i, column=6, value=f'=COUNTIFS(Calendar!$C$5:$C${last},D{i},Calendar!$H$5:$H${last},"Published")')
    sm["A12"] = "Published this month"
    sm["A12"].font = Font(bold=True)
    sm["B12"] = (f'=COUNTIFS(Calendar!$H$5:$H${last},"Published",Calendar!$A$5:$A${last},">="&DATE(YEAR(TODAY()),MONTH(TODAY()),1),'
                 f'Calendar!$A$5:$A${last},"<"&DATE(YEAR(TODAY()),MONTH(TODAY())+1,1))')
    sm["A13"] = "Pillar mix"
    sm["A13"].font = Font(bold=True, color=BRAND)
    for i, p in enumerate(PILLARS, start=14):
        sm.cell(row=i, column=1, value=p)
        sm.cell(row=i, column=2, value=f'=COUNTIF(Calendar!$D$5:$D${last},A{i})')
        sm.cell(row=i, column=3, value=f'=IF(SUM($B$14:$B$18)=0,"",B{i}/SUM($B$14:$B$18))').number_format = "0%"
    wb.active = 0
    wb.save(path)


# --------------------------------------------------------------------------
# 3. Client CRM & Pipeline
# --------------------------------------------------------------------------
STAGES = [("Lead", 0.05), ("Contacted", 0.10), ("Discovery call", 0.25), ("Proposal sent", 0.50),
          ("Negotiation", 0.75), ("Won", 1.0), ("Lost", 0.0)]
SOURCES = ["Referral", "LinkedIn", "Instagram", "Cold email", "Website", "Marketplace", "Event", "Other"]
CRM_ROWS = 500


def build_crm(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Pipeline"
    _title(ws, "Client CRM & Sales Pipeline", "Yellow = due today or earlier. Probability and weighted value fill in from the stage.")
    headers = ["Name", "Company", "Email", "Source", "Stage", "Deal value", "Probability",
               "Weighted value", "Last contact", "Next follow-up", "Next action", "Notes"]
    _header(ws, 4, headers, [20, 20, 26, 13, 16, 13, 11, 14, 13, 14, 30, 30])
    last = 4 + CRM_ROWS

    stg = wb.create_sheet("Stages")
    _title(stg, "Stages & Probabilities", "Edit win probabilities to match your history.")
    _header(stg, 4, ["Stage", "Win probability"], [18, 16])
    stg.freeze_panes = None
    for i, (s, p) in enumerate(STAGES, start=5):
        stg.cell(row=i, column=1, value=s)
        c = stg.cell(row=i, column=2, value=p)
        c.number_format = "0%"
        c.fill = INPUT_FILL
    stg["D4"] = "Sources"
    stg["D4"].font = Font(bold=True)
    for i, s in enumerate(SOURCES, start=5):
        stg.cell(row=i, column=4, value=s)
    stg.column_dimensions["D"].width = 16

    _list_validation(ws, "=Stages!$A$5:$A$11", f"E5:E{last}")
    _list_validation(ws, "=Stages!$D$5:$D$15", f"D5:D{last}")
    for r in range(5, last + 1):
        ws.cell(row=r, column=6).number_format = MONEY
        ws.cell(row=r, column=7, value=f'=IF(E{r}="","",IFERROR(VLOOKUP(E{r},Stages!$A$5:$B$11,2,FALSE),""))').number_format = "0%"
        ws.cell(row=r, column=8, value=f'=IF(OR(F{r}="",G{r}=""),"",F{r}*G{r})').number_format = MONEY
        ws.cell(row=r, column=9).number_format = "yyyy-mm-dd"
        ws.cell(row=r, column=10).number_format = "yyyy-mm-dd"
    ws.conditional_formatting.add(
        f"J5:J{last}",
        FormulaRule(formula=[f'AND(J5<>"",J5<=TODAY(),$E5<>"Won",$E5<>"Lost")'], fill=INPUT_FILL, font=Font(bold=True, color="B91C1C")))
    ws.conditional_formatting.add(f"A5:L{last}", FormulaRule(formula=['$E5="Won"'], fill=GREEN_FILL))
    ws.conditional_formatting.add(f"A5:L{last}", FormulaRule(formula=['$E5="Lost"'], font=Font(color="9CA3AF")))
    ws.auto_filter.ref = f"A4:L{last}"
    ws["A5"], ws["B5"], ws["C5"], ws["D5"], ws["E5"], ws["F5"] = (
        "Sample Lead", "Example Ltd", "hello@example.com", "Referral", "Proposal sent", 2500)
    ws["K5"] = "Follow up on proposal. Delete this sample row."

    fc = wb.create_sheet("Forecast", 0)
    _title(fc, "Pipeline Forecast", "Updates automatically from the Pipeline tab.")
    _header(fc, 4, ["Stage", "Deals", "Deal value", "Weighted value"], [18, 9, 15, 16])
    fc.freeze_panes = None
    for i, (s, _) in enumerate(STAGES, start=5):
        fc.cell(row=i, column=1, value=s)
        fc.cell(row=i, column=2, value=f"=COUNTIF(Pipeline!$E$5:$E${last},A{i})")
        fc.cell(row=i, column=3, value=f"=SUMIF(Pipeline!$E$5:$E${last},A{i},Pipeline!$F$5:$F${last})").number_format = MONEY
        fc.cell(row=i, column=4, value=f"=SUMIF(Pipeline!$E$5:$E${last},A{i},Pipeline!$H$5:$H${last})").number_format = MONEY
    fc["A13"] = "Open pipeline"
    fc["B13"] = "=SUM(B5:B9)"
    fc["C13"] = "=SUM(C5:C9)"
    fc["D13"] = "=SUM(D5:D9)"
    fc["A14"] = "Win rate (won / closed)"
    fc["B14"] = '=IF(B10+B11=0,"",B10/(B10+B11))'
    fc["B14"].number_format = "0%"
    fc["A15"] = "Follow-ups due"
    fc["B15"] = (f'=COUNTIFS(Pipeline!$J$5:$J${last},"<="&TODAY(),Pipeline!$E$5:$E${last},"<>Won",'
                 f'Pipeline!$E$5:$E${last},"<>Lost")')
    for r in (13, 14, 15):
        for c in range(1, 5):
            fc.cell(row=r, column=c).fill = TOTAL_FILL
            fc.cell(row=r, column=c).font = Font(bold=True)
    fc["C13"].number_format = fc["D13"].number_format = MONEY
    wb.active = 0
    wb.save(path)
